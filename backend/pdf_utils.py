from reportlab.lib.pagesizes import A4, letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
    PageBreak, Image, KeepTogether
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT, TA_JUSTIFY
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.piecharts import Pie
from datetime import datetime
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

class PDFGenerator:
    """Enhanced PDF Generator for professional variation proposals"""
    
    @staticmethod
    def _safe_float(value, default=0.0):
        try:
            if value is None or value == "":
                return default
            return float(str(value).replace(",", "").replace("Rs.", "").replace("Rs", "").strip())
        except Exception:
            return default

    @staticmethod
    def _money(value):
        return f"Rs. {PDFGenerator._safe_float(value):,.2f}"

    @staticmethod
    def _confidence_text(value):
        if value is None or value == "":
            return ""
        try:
            return f"{float(value):.2f}"
        except Exception:
            return str(value)
        
    @staticmethod
    def _create_header(elements, styles, project_name="Construction Project"):
        """Create professional header"""
        header_style = ParagraphStyle(
            'HeaderStyle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#1a365d'),
            alignment=TA_CENTER,
            spaceAfter=6,
            fontName='Helvetica-Bold'
        )
        
        subheader_style = ParagraphStyle(
            'SubHeaderStyle',
            parent=styles['Normal'],
            fontSize=14,
            textColor=colors.HexColor('#2d3748'),
            alignment=TA_CENTER,
            spaceAfter=20
        )
        
        elements.append(Paragraph("VARIATION PROPOSAL", header_style))
        elements.append(Paragraph(f"Project: {project_name}", subheader_style))
        elements.append(Spacer(1, 0.2*inch))
    
    @staticmethod
    def _create_disclaimer_section(elements, styles):
        """Create disclaimer for ML-assisted extraction and rule-based evaluation"""
        disclaimer_style = ParagraphStyle(
            'DisclaimerStyle',
            parent=styles['Normal'],
            fontSize=9,
            textColor=colors.HexColor('#742a2a'),
            alignment=TA_JUSTIFY,
            spaceAfter=12,
            borderWidth=1,
            borderColor=colors.HexColor('#fc8181'),
            borderPadding=8,
            backColor=colors.HexColor('#fed7d7')
        )
        
        disclaimer_text = (
            "<b>IMPORTANT DISCLAIMER:</b> ML-assisted extraction was used to structure and parse input data from "
            "supporting documents. Final cost and time impact were calculated using deterministic rule-based logic "
            "with evidence-based rate and productivity sources. This proposal requires professional QS/Engineer "
            "verification before formal submission to the Engineer/Employer. The System provides a calculation "
            "framework; professional judgment is essential for final approval."
        )
        
        elements.append(Paragraph(disclaimer_text, disclaimer_style))
        elements.append(Spacer(1, 0.1*inch))
    
    @staticmethod
    def _create_info_section(elements, styles, proposal_data):
        """Create proposal information section"""
        info_style = ParagraphStyle(
            'InfoStyle',
            parent=styles['Normal'],
            fontSize=10,
            leading=14
        )
        
        curr_date = datetime.now().strftime("%d %B %Y")
        variation_id = proposal_data.get('variation_id', 'DRAFT')
        variation_type = proposal_data.get('variation_type', 'TYPE1')
        evaluation_mode = proposal_data.get('evaluation_mode', 'quantity_change')
        human_confirmed = "Yes" if proposal_data.get('human_confirmed') else "No"
        
        info_data = [
            ["Proposal Reference:", f"VAR-{variation_id}"],
            ["Date of Submission:", curr_date],
            ["Variation Type:", variation_type],
            ["Evaluation Mode:", evaluation_mode],
            ["Human Confirmed:", human_confirmed],
            ["Status:", proposal_data.get('status', 'Under Review')]
        ]
        
        info_table = Table(info_data, colWidths=[2.5*inch, 4*inch])
        info_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTNAME', (1, 0), (1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#2d3748')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ]))
        
        elements.append(info_table)
        elements.append(Spacer(1, 0.3*inch))
    
    @staticmethod
    def _create_section_header(elements, styles, title):
        """Create section header"""
        section_style = ParagraphStyle(
            'SectionStyle',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor('#1a365d'),
            spaceAfter=12,
            spaceBefore=12,
            fontName='Helvetica-Bold',
            borderWidth=0,
            borderColor=colors.HexColor('#cbd5e0'),
            borderPadding=8,
            backColor=colors.HexColor('#edf2f7')
        )
        
        elements.append(Paragraph(title, section_style))
    
    @staticmethod
    def _create_description_section(elements, styles, proposal_data):
        """Create variation description section"""
        PDFGenerator._create_section_header(elements, styles, "1. VARIATION DESCRIPTION")
        
        description = proposal_data.get('description', 'No description provided')
        
        desc_style = ParagraphStyle(
            'DescStyle',
            parent=styles['Normal'],
            fontSize=11,
            leading=16,
            alignment=TA_JUSTIFY
        )
        
        elements.append(Paragraph(description, desc_style))
        elements.append(Spacer(1, 0.2*inch))
    
    @staticmethod
    def _create_cost_breakdown_section(elements, styles, proposal_data):
        """Create detailed cost breakdown section"""
        PDFGenerator._create_section_header(elements, styles, "2. COST IMPACT ANALYSIS")

        small_style = ParagraphStyle(
            'SmallTableText',
            parent=styles['Normal'],
            fontSize=7,
            leading=9
        )

        cost_lines = proposal_data.get('cost_lines', []) or []
        table_data = [["Type", "Description", "Qty", "Unit", "Rate", "Amount", "Formula"]]
        total_impact = 0.0

        for line in cost_lines:
            if not isinstance(line, dict):
                continue

            amount = PDFGenerator._safe_float(line.get('amount', 0))
            total_impact += amount

            table_data.append([
                str(line.get('line_type', 'n/a')),
                Paragraph(str(line.get('description', 'N/A'))[:85], small_style),
                f"{PDFGenerator._safe_float(line.get('quantity', 0)):.2f}",
                str(line.get('unit', '')),
                PDFGenerator._money(line.get('rate', 0)),
                PDFGenerator._money(amount),
                Paragraph(str(line.get('formula', ''))[:80], small_style),
            ])

        if not cost_lines:
            table_data.append([
                "manual",
                Paragraph("No confirmed cost line available. QS review required.", small_style),
                "", "", "", "", ""
            ])

        table_data.append([
            Paragraph("<b>TOTAL</b>", small_style),
            "", "", "", "",
            Paragraph(f"<b>{PDFGenerator._money(total_impact)}</b>", small_style),
            ""
        ])

        # A4 usable width is limited. Keep total width below page frame.
        cost_table = Table(
            table_data,
            colWidths=[
                0.65*inch,
                1.65*inch,
                0.55*inch,
                0.40*inch,
                0.85*inch,
                0.90*inch,
                1.65*inch
            ],
            repeatRows=1
        )

        cost_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2c5282')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 7),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),

            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 7),
            ('ALIGN', (2, 1), (-1, -1), 'CENTER'),
            ('ALIGN', (1, 1), (1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),

            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#edf2f7')),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),

            ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#2c5282')),
        ]))

        elements.append(cost_table)
        elements.append(Spacer(1, 0.25*inch))
    
    @staticmethod
    def _create_time_impact_section(elements, styles, proposal_data):
        """Create time impact section with honest CPM/EOT wording"""
        time_impact = proposal_data.get('time_impact') or proposal_data.get('time_impact_result')

        if not time_impact:
            return

        PDFGenerator._create_section_header(elements, styles, "3. TIME IMPACT ANALYSIS")

        manual_required = bool(time_impact.get("manual_required"))
        cpm_proof = bool(time_impact.get("cpm_proof"))
        additional_duration = PDFGenerator._safe_float(time_impact.get("additional_duration", 0))
        eot_days = PDFGenerator._safe_float(time_impact.get("eot_days", 0))

        if manual_required:
            summary_text = (
                "<b>Time impact cannot be finalised.</b> "
                f"{time_impact.get('message', 'Activity mapping and/or productivity evidence is missing.')}"
            )
        elif not cpm_proof:
            summary_text = (
                f"The additional duration is calculated as <b>{additional_duration:.2f} days</b>. "
                "However, EOT is not finalised because CPM activity proof is not available. "
                "Planner/QS review is required before claiming EOT."
            )
        else:
            summary_text = (
                f"The proposed variation results in an additional duration of <b>{additional_duration:.2f} days</b>. "
                f"The Extension of Time (EOT) assessed through CPM logic is <b>{eot_days:.2f} days</b>."
            )

        elements.append(Paragraph(summary_text, styles['Normal']))
        elements.append(Spacer(1, 0.1*inch))

        eot_data = [
            ["Parameter", "Value"],
            ["Affected Activity", str(time_impact.get('activity_name') or time_impact.get('activity_ref') or 'N/A')],
            ["Activity Reference", str(time_impact.get('activity_ref', ''))],
            ["Productivity", str(time_impact.get('productivity', ''))],
            ["Productivity Source", str(time_impact.get('productivity_source', ''))],
            ["Additional Duration", f"{additional_duration:.2f} days"],
            ["Formula", str(time_impact.get('formula', ''))],
            ["On Critical Path", "Yes" if time_impact.get('is_critical') else "No"],
            ["Original Float", f"{PDFGenerator._safe_float(time_impact.get('original_float', 0)):.2f} days"],
            ["CPM Proof Available", "Yes" if cpm_proof else "No"],
            ["Extension of Time (EOT)", f"{eot_days:.2f} days" if cpm_proof else "Not finalised"],
        ]

        eot_table = Table(eot_data, colWidths=[2.3*inch, 4.2*inch], repeatRows=1)
        eot_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2c5282')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTNAME', (0, 1), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f7fafc')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#2c5282')),
        ]))

        elements.append(eot_table)
        elements.append(Spacer(1, 0.2*inch))

    @staticmethod
    def _create_evidence_section(elements, styles, proposal_data):
        """Show only rate evidence actually used in calculation"""
        PDFGenerator._create_section_header(elements, styles, "4. RATE SOURCE EVIDENCE USED IN CALCULATION")

        small_style = ParagraphStyle(
            'EvidenceSmallText',
            parent=styles['Normal'],
            fontSize=7,
            leading=9
        )

        sources = proposal_data.get('used_rate_sources') or proposal_data.get('rate_sources') or []

        evidence_rows = [["Source", "Item Ref", "Description", "Unit", "Rate", "Confidence"]]

        for source in sources:
            if not isinstance(source, dict):
                continue

            evidence_rows.append([
                str(source.get('source_type', '')),
                str(source.get('item_reference', '')),
                Paragraph(str(source.get('description', ''))[:90], small_style),
                str(source.get('unit', '')),
                PDFGenerator._money(source.get('rate', 0)),
                PDFGenerator._confidence_text(source.get('confidence', source.get('confidence_score', ''))),
            ])

        if len(evidence_rows) == 1:
            evidence_rows.append(["n/a", "", "No used rate evidence supplied", "", "", ""])

        evidence_table = Table(
            evidence_rows,
            colWidths=[
                0.9*inch,
                0.9*inch,
                2.5*inch,
                0.5*inch,
                1.0*inch,
                1.0*inch
            ],
            repeatRows=1
        )

        evidence_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2c5282')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 7),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 7),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f7fafc')]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#2c5282')),
        ]))

        elements.append(evidence_table)
        elements.append(Spacer(1, 0.2*inch))
    
    @staticmethod
    def _create_validation_section(elements, styles, proposal_data):
        """Create QS validation results section"""
        validation = proposal_data.get('validation') or proposal_data.get('validation_results')

        if not validation:
            return

        PDFGenerator._create_section_header(elements, styles, "5. QS VALIDATION RESULTS")

        warnings = validation.get('warnings', []) if isinstance(validation, dict) else []
        errors = validation.get('errors', []) if isinstance(validation, dict) else []

        if errors:
            status = "FAILED"
            status_color = colors.HexColor('#e53e3e')
        elif warnings:
            status = "WARNINGS"
            status_color = colors.HexColor('#d69e2e')
        else:
            status = "PASSED"
            status_color = colors.HexColor('#38a169')

        status_text = f"<font color='{status_color.hexval()}'>●</font> <b>Status: {status}</b>"
        elements.append(Paragraph(status_text, styles['Normal']))
        elements.append(Spacer(1, 0.1*inch))

        if warnings:
            elements.append(Paragraph("<b>Warnings:</b>", styles['Heading4']))
            for warning in warnings:
                elements.append(Paragraph(f"- {warning}", styles['Normal']))
            elements.append(Spacer(1, 0.1*inch))

        if errors:
            elements.append(Paragraph("<b>Errors:</b>", styles['Heading4']))
            for error in errors:
                elements.append(Paragraph(f"- {error}", styles['Normal']))
            elements.append(Spacer(1, 0.1*inch))
    
    @staticmethod
    def _create_summary_section(elements, styles, proposal_data):
        """Create executive summary"""
        PDFGenerator._create_section_header(elements, styles, "6. EXECUTIVE SUMMARY")
        
        total_cost = float(proposal_data.get('total_cost_impact', proposal_data.get('cost_impact', 0)) or 0)
        time_impact = proposal_data.get('time_impact') or {}
        total_time = float(time_impact.get('eot_days', proposal_data.get('time_impact_days', 0)) or 0)
        
        summary_data = [
            ["Total Cost Impact:", f"Rs. {total_cost:,.2f}"],
            ["Total Time Impact:", f"{total_time:.1f} days"],
            ["Recommendation:", proposal_data.get('recommendation', 'Approve subject to validation')]
        ]
        
        summary_table = Table(summary_data, colWidths=[2.5*inch, 4*inch])
        summary_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTNAME', (1, 0), (1, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 12),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#1a365d')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#edf2f7')),
            ('BOX', (0, 0), (-1, -1), 2, colors.HexColor('#2c5282')),
        ]))
        
        elements.append(summary_table)
        elements.append(Spacer(1, 0.3*inch))
    
    @staticmethod
    def _create_footer(elements, styles):
        """Create signature section"""
        elements.append(Spacer(1, 0.5*inch))
        
        footer_style = ParagraphStyle(
            'FooterStyle',
            parent=styles['Normal'],
            fontSize=10,
            alignment=TA_CENTER
        )
        
        elements.append(Paragraph("_" * 40, footer_style))
        elements.append(Paragraph("QS / Engineer Signature", footer_style))
        elements.append(Spacer(1, 0.1*inch))
        elements.append(Paragraph(f"Generated: {datetime.now().strftime('%d %B %Y at %H:%M')}", footer_style))
    
    @staticmethod
    def generate_variation_proposal(proposal_data, output_path):
        """
        Generate comprehensive variation proposal PDF
        
        Args:
            proposal_data: Dictionary containing all proposal information
            output_path: Path where PDF will be saved
            
        Returns:
            Path to generated PDF
        """
        output_dir = os.path.join("backend", "generated_reports")
        os.makedirs(output_dir, exist_ok=True)
        if not output_path:
            output_path = os.path.join(output_dir, f"variation_proposal_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")
        elif not os.path.isabs(output_path):
            output_path = os.path.join(output_dir, os.path.basename(output_path))

        doc = SimpleDocTemplate(
            output_path,
            pagesize=A4,
            rightMargin=0.75*inch,
            leftMargin=0.75*inch,
            topMargin=0.75*inch,
            bottomMargin=0.75*inch
        )
        
        styles = getSampleStyleSheet()
        elements = []
        
        # Build document sections
        project_name = proposal_data.get('project_name', 'Construction Project')
        
        PDFGenerator._create_header(elements, styles, project_name)
        PDFGenerator._create_disclaimer_section(elements, styles)
        PDFGenerator._create_info_section(elements, styles, proposal_data)
        PDFGenerator._create_description_section(elements, styles, proposal_data)
        PDFGenerator._create_cost_breakdown_section(elements, styles, proposal_data)
        PDFGenerator._create_time_impact_section(elements, styles, proposal_data)
        PDFGenerator._create_evidence_section(elements, styles, proposal_data)
        PDFGenerator._create_validation_section(elements, styles, proposal_data)
        PDFGenerator._create_summary_section(elements, styles, proposal_data)
        PDFGenerator._create_footer(elements, styles)
        
        # Build PDF
        doc.build(elements)
        
        return output_path

    @staticmethod
    def generate_variation_docx(proposal_data, output_path=None):
        """
        Generate an editable DOCX variation proposal with only used evidence.
        Returns path to the created DOCX file.
        """
        output_dir = os.path.join("backend", "generated_reports")
        os.makedirs(output_dir, exist_ok=True)
        if not output_path:
            output_path = os.path.join(output_dir, f"variation_proposal_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
        elif not os.path.isabs(output_path):
            output_path = os.path.join(output_dir, os.path.basename(output_path))

        doc = Document()
        style = doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)

        # Header
        h = doc.add_heading('VARIATION PROPOSAL', level=1)
        h.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        doc.add_paragraph(f"Project: {proposal_data.get('project_name', '')}").alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        doc.add_paragraph('')

        # Info table
        info = doc.add_table(rows=0, cols=2)
        info.style = 'Light List'
        info.add_row().cells[0].text = 'Proposal Reference:'
        info.add_row().cells[0].text = ''  # placeholder to keep structure
        # Use a simple approach: add key lines
        info.add_row().cells[0].text = 'Variation Reference:'
        info.add_row().cells[1].text = str(proposal_data.get('variation_id', 'DRAFT'))
        info.add_row().cells[0].text = 'Variation Type:'
        info.add_row().cells[1].text = str(proposal_data.get('variation_type', ''))
        info.add_row().cells[0].text = 'Evaluation Mode:'
        info.add_row().cells[1].text = str(proposal_data.get('evaluation_mode', ''))
        doc.add_paragraph('')

        # Description
        doc.add_heading('1. Variation Description', level=2)
        doc.add_paragraph(proposal_data.get('description', ''))

        # Human confirmation summary
        doc.add_heading('2. Human Confirmation Summary', level=2)
        confirmed = 'Yes' if proposal_data.get('human_confirmed') else 'No'
        doc.add_paragraph(f"Human Confirmed: {confirmed}")
        doc.add_paragraph('Confirmed Values:')
        for k in ['variation_type','evaluation_mode']:
            doc.add_paragraph(f"- {k}: {proposal_data.get(k)}")

        # Cost Impact table
        doc.add_heading('3. Cost Impact (Cost Lines)', level=2)
        table = doc.add_table(rows=1, cols=6)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Line ID'
        hdr_cells[1].text = 'Description'
        hdr_cells[2].text = 'Qty'
        hdr_cells[3].text = 'Unit'
        hdr_cells[4].text = 'Rate'
        hdr_cells[5].text = 'Amount'
        for line in proposal_data.get('cost_lines', []) or []:
            row = table.add_row().cells
            row[0].text = str(line.get('id', line.get('line_id', '')))
            row[1].text = str(line.get('description', ''))
            row[2].text = str(line.get('quantity', ''))
            row[3].text = str(line.get('unit', ''))
            row[4].text = f"Rs. {float(line.get('rate', 0) or 0):,.2f}"
            row[5].text = f"Rs. {float(line.get('amount', 0) or 0):,.2f}"

        # Time impact table
        doc.add_heading('4. Time Impact', level=2)
        time_impact = proposal_data.get('time_impact') or {}
        if not time_impact or not time_impact.get('activity_name') or not time_impact.get('additional_duration'):
            doc.add_paragraph('Time impact cannot be finalised because activity mapping/productivity is missing.')
        else:
            doc.add_paragraph(f"Affected Activity: {time_impact.get('activity_name')}")
            doc.add_paragraph(f"Original Duration: {time_impact.get('original_duration', '')} days")
            doc.add_paragraph(f"Additional Duration: {time_impact.get('additional_duration', '')} days")
            doc.add_paragraph(f"Revised Duration: {time_impact.get('revised_duration', '')} days")
            doc.add_paragraph(f"On Critical Path: {'Yes' if time_impact.get('is_critical') else 'No'}")

        # Rate source evidence (only used)
        doc.add_heading('5. Rate Source Evidence (Used)', level=2)
        for src in proposal_data.get('rate_sources', []) or []:
            doc.add_paragraph(f"- Source Type: {src.get('source_type')}; File: {src.get('source_file')}; Item Ref: {src.get('item_reference')}; Rate: Rs. {src.get('rate')}; Confidence: {src.get('confidence')}")

        # Productivity/activity evidence
        doc.add_heading('6. Productivity / Activity Evidence (Used)', level=2)
        prod_evidence = proposal_data.get('productivity_evidence', []) or []
        if not prod_evidence:
            doc.add_paragraph('No productivity evidence was used in the calculation.')
        else:
            for pe in prod_evidence:
                doc.add_paragraph(f"- {pe.get('source_type')} | {pe.get('source_file')} | {pe.get('description')} | {pe.get('productivity')}")

        # Validation warnings/errors
        doc.add_heading('7. Validation Warnings / Errors', level=2)
        validation = proposal_data.get('validation') or {}
        for warn in validation.get('warnings', []):
            doc.add_paragraph(f"- WARNING: {warn}")
        for err in validation.get('errors', []):
            doc.add_paragraph(f"- ERROR: {err}")

        # Assumptions
        doc.add_heading('8. Assumptions', level=2)
        for a in proposal_data.get('assumptions', []):
            doc.add_paragraph(f"- {a}")

        # Editable notes
        doc.add_heading('9. QS / Engineer Notes (Editable)', level=2)
        doc.add_paragraph('\n' * 6)

        # Signature
        doc.add_heading('10. Signatures', level=2)
        doc.add_paragraph('QS: ______________________        Date: ____________')
        doc.add_paragraph('Engineer: ___________________     Date: ____________')

        doc.save(output_path)
        return output_path
