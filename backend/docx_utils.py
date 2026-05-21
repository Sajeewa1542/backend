from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
import os

TABLE_GRID_STYLE = "Table Grid"
LIST_BULLET_STYLE = "List Bullet"


class DOCXGenerator:
    """Editable Word variation proposal generator for QS review."""

    @staticmethod
    def _money(value):
        try:
            return f"Rs. {float(value):,.2f}"
        except Exception:
            return "Rs. 0.00"

    @staticmethod
    def _num(value):
        try:
            return f"{float(value):,.2f}"
        except Exception:
            return ""

    @staticmethod
    def _add_heading(doc, text, level=1):
        doc.add_heading(text, level=level)

    @staticmethod
    def _add_kv_table(doc, rows):
        table = doc.add_table(rows=1, cols=2)
        table.style = TABLE_GRID_STYLE
        hdr = table.rows[0].cells
        hdr[0].text = "Item"
        hdr[1].text = "Details"

        for key, value in rows:
            cells = table.add_row().cells
            cells[0].text = str(key)
            cells[1].text = "" if value is None else str(value)

        doc.add_paragraph("")
        return table

    @staticmethod
    def generate_variation_proposal(proposal_data, output_path):
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        doc = Document()

        section = doc.sections[0]
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("VARIATION PROPOSAL")
        run.bold = True
        run.font.size = Pt(18)

        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.add_run(f"Project: {proposal_data.get('project_name', 'Construction Project')}").bold = True

        doc.add_paragraph("")

        disclaimer = doc.add_paragraph()
        disclaimer.add_run("IMPORTANT DISCLAIMER: ").bold = True
        disclaimer.add_run(
            "ML-assisted extraction was used only to structure input data and identify candidate evidence. "
            "Final cost and time impacts are calculated using deterministic rule-based logic after human confirmation. "
            "This editable Word proposal requires QS/Engineer review before formal submission."
        )

        DOCXGenerator._add_heading(doc, "1. Proposal Information", 1)
        DOCXGenerator._add_kv_table(doc, [
            ("Proposal Reference", f"VAR-{proposal_data.get('variation_id', 'DRAFT')}"),
            ("Date", datetime.now().strftime("%d %B %Y")),
            ("Variation Type", proposal_data.get("variation_type", "")),
            ("Evaluation Mode", proposal_data.get("evaluation_mode", "")),
            ("Human Confirmed", "Yes" if proposal_data.get("human_confirmed") else "No"),
            ("Status", proposal_data.get("status", "Draft")),
        ])

        DOCXGenerator._add_heading(doc, "2. Variation Description", 1)
        doc.add_paragraph(proposal_data.get("variation_description") or proposal_data.get("description") or "")

        DOCXGenerator._add_heading(doc, "3. Human Confirmation Summary", 1)
        DOCXGenerator._add_kv_table(doc, [
            ("Original BOQ Item", proposal_data.get("original_boq_item_ref") or proposal_data.get("confirmed_boq_item_id", "")),
            ("Original Description", proposal_data.get("original_description", "")),
            ("Original Quantity", proposal_data.get("original_quantity", "")),
            ("New / Replacement Quantity", proposal_data.get("new_quantity") or proposal_data.get("replacement_quantity", "")),
            ("Unit", proposal_data.get("unit", "")),
            ("Confirmed Rate", DOCXGenerator._money(proposal_data.get("confirmed_rate", 0))),
            ("Rate Source", proposal_data.get("confirmed_rate_source", "")),
            ("Activity", proposal_data.get("confirmed_activity_ref") or proposal_data.get("confirmed_activity_id", "")),
            ("Productivity", proposal_data.get("confirmed_productivity", "")),
            ("Productivity Source", proposal_data.get("productivity_source", "")),
        ])

        DOCXGenerator._add_heading(doc, "4. Cost Impact Analysis", 1)

        cost_lines = proposal_data.get("cost_lines", [])
        table = doc.add_table(rows=1, cols=7)
        table.style = TABLE_GRID_STYLE
        headers = ["Line Type", "Description", "Qty", "Unit", "Rate", "Amount", "Formula"]
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h

        if cost_lines:
            for line in cost_lines:
                cells = table.add_row().cells
                cells[0].text = str(line.get("line_type", ""))
                cells[1].text = str(line.get("description", ""))
                cells[2].text = DOCXGenerator._num(line.get("quantity", 0))
                cells[3].text = str(line.get("unit", ""))
                cells[4].text = DOCXGenerator._money(line.get("rate", 0))
                cells[5].text = DOCXGenerator._money(line.get("amount", 0))
                cells[6].text = str(line.get("formula", ""))
        else:
            cells = table.add_row().cells
            cells[0].text = "No cost lines generated"
            cells[1].text = "Manual QS review required"

        doc.add_paragraph("")
        p = doc.add_paragraph()
        p.add_run("Total Cost Impact: ").bold = True
        p.add_run(DOCXGenerator._money(proposal_data.get("total_cost_impact", proposal_data.get("cost_impact", 0))))

        DOCXGenerator._add_heading(doc, "5. Time Impact Analysis", 1)
        time_data = proposal_data.get("time_impact_result") or proposal_data.get("time_impact", {})

        if isinstance(time_data, dict) and time_data.get("manual_required"):
            doc.add_paragraph(
                "Time impact cannot be finalised because activity mapping and/or productivity evidence is missing."
            )
        elif isinstance(time_data, dict):
            DOCXGenerator._add_kv_table(doc, [
                ("Activity Reference", time_data.get("activity_ref", "")),
                ("Activity Name", time_data.get("activity_name", "")),
                ("Productivity", time_data.get("productivity", "")),
                ("Productivity Source", time_data.get("productivity_source", "")),
                ("Additional Duration", f"{time_data.get('additional_duration', 0)} days"),
                ("Critical Path", "Yes" if time_data.get("is_critical") else "No"),
                ("Float", f"{time_data.get('float', 0)} days"),
                ("EOT", f"{time_data.get('eot_days', 0)} days"),
                ("Formula", time_data.get("formula", "")),
            ])
        else:
            doc.add_paragraph("Time impact information is not available. Manual QS/Planner review required.")

        DOCXGenerator._add_heading(doc, "6. Rate Source Evidence Used", 1)
        used_sources = proposal_data.get("used_rate_sources", [])
        table = doc.add_table(rows=1, cols=7)
        table.style = TABLE_GRID_STYLE
        headers = ["Source Type", "Source File", "Reference", "Description", "Unit", "Rate", "Confidence"]
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h

        if used_sources:
            for src in used_sources:
                cells = table.add_row().cells
                cells[0].text = str(src.get("source_type", ""))
                cells[1].text = str(src.get("source_file", ""))
                cells[2].text = str(src.get("item_reference", ""))
                cells[3].text = str(src.get("description", ""))
                cells[4].text = str(src.get("unit", ""))
                cells[5].text = DOCXGenerator._money(src.get("rate", 0))
                cells[6].text = str(src.get("confidence", ""))
        else:
            cells = table.add_row().cells
            cells[0].text = "No external source evidence used"
            cells[1].text = "BOQ/manual source or QS review required"

        DOCXGenerator._add_heading(doc, "7. Validation Checklist", 1)
        validation = proposal_data.get("validation", {}) or proposal_data.get("validation_results", {})
        warnings = validation.get("warnings", []) if isinstance(validation, dict) else []
        errors = validation.get("errors", []) if isinstance(validation, dict) else []

        if not warnings and not errors:
            doc.add_paragraph("No major validation warnings recorded.")
        else:
            if warnings:
                doc.add_paragraph("Warnings:")
                for w in warnings:
                    doc.add_paragraph(str(w), style=LIST_BULLET_STYLE)
            if errors:
                doc.add_paragraph("Errors:")
                for e in errors:
                    doc.add_paragraph(str(e), style=LIST_BULLET_STYLE)

        DOCXGenerator._add_heading(doc, "8. Assumptions and QS Editable Notes", 1)
        doc.add_paragraph("Assumptions:")
        for a in proposal_data.get("assumptions", ["Rates and productivity values are subject to QS/Engineer review."]):
            doc.add_paragraph(str(a), style=LIST_BULLET_STYLE)

        doc.add_paragraph("")
        doc.add_paragraph("QS / Engineer Notes:")
        for _ in range(5):
            doc.add_paragraph("........................................................................................................")

        DOCXGenerator._add_heading(doc, "9. Signature", 1)
        doc.add_paragraph("Prepared by: ....................................................")
        doc.add_paragraph("Checked by QS/Engineer: ..........................................")
        doc.add_paragraph("Date: ............................................................")

        doc.save(output_path)
        return output_path